#!/usr/bin/env python3
"""
Import CSMC GIS Excel exports → public/data/*.json (FacilityRecord shape).

Source files live in data/gis-source/. Re-run after GIS updates:
  python scripts/import-gis-facilities.py

Does not invent coordinates, phones, or names. Empty source fields stay empty.
"""
from __future__ import annotations

import json
import re
from pathlib import Path

try:
    import openpyxl
except ImportError as e:
    raise SystemExit("openpyxl is required: pip install openpyxl") from e

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "data" / "gis-source"
OUT = ROOT / "public" / "data"
REPORT = ROOT / "data" / "gis-import-report.json"


def cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def slugify(text: str, fallback: str) -> str:
    s = re.sub(r"[^a-zA-Z0-9]+", "-", (text or "").lower()).strip("-")
    return s[:48] or fallback


def maps_url(lat: str, lng: str, name: str, address: str) -> str:
    if lat and lng:
        return f"https://www.google.com/maps?q={lat},{lng}"
    q = " ".join(p for p in [name, address, "Chhatrapati Sambhajinagar"] if p)
    if not q:
        return ""
    from urllib.parse import quote_plus

    return f"https://www.google.com/maps/search/?api=1&query={quote_plus(q)}"


def parse_coord(raw: str) -> str:
    raw = (raw or "").strip()
    if not raw or raw in {"0", "0.0", "0.00"}:
        return ""
    try:
        v = float(raw)
    except ValueError:
        return ""
    if abs(v) < 0.1:
        return ""
    # Keep source precision without fabricating digits
    return raw if re.search(r"\.", raw) else str(v)


def read_sheet(path: Path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    headers = [cell(c) or f"COL_{i}" for i, c in enumerate(rows[0])]
    data = []
    empty = 0
    for r in rows[1:]:
        vals = [cell(c) for c in r]
        if not any(vals):
            empty += 1
            continue
        # pad/truncate to headers
        while len(vals) < len(headers):
            vals.append("")
        data.append(dict(zip(headers, vals[: len(headers)])))
    return data, empty


def detail(label_en: str, label_mr: str, value: str):
    value = (value or "").strip()
    if not value or value in {"0", "None", "-"}:
        return None
    return {"labelEn": label_en, "labelMr": label_mr, "value": value}


def facility(
    *,
    id_: str,
    name: str,
    address: str = "",
    zone: str = "",
    phone: str = "",
    timings: str = "",
    latitude: str = "",
    longitude: str = "",
    details: list | None = None,
):
    lat = parse_coord(latitude)
    lng = parse_coord(longitude)
    # Both required for a usable point
    if not lat or not lng:
        lat, lng = "", ""
    details = [d for d in (details or []) if d]
    rec = {
        "id": id_,
        "name": name.strip(),
        "address": address.strip(),
        "phone": phone.strip(),
        "timings": timings.strip(),
        "latitude": lat,
        "longitude": lng,
        "googleMapsUrl": maps_url(lat, lng, name, address),
    }
    if zone.strip():
        rec["zone"] = zone.strip()
    if details:
        rec["details"] = details
    return rec


def import_schools():
    rows, empty = read_sheet(SRC / "CSMC_School.xlsx")
    out = []
    for i, r in enumerate(rows, 1):
        name = r.get("Name", "")
        if not name:
            continue
        zone = r.get("Zone", "")
        out.append(
            facility(
                id_=f"school-{i:03d}",
                name=name,
                address=r.get("Address", ""),
                zone=f"Zone {zone}" if zone else "",
                details=[detail("Zone", "झोन", f"Zone {zone}" if zone else "")],
            )
        )
    return out, {"sourceRows": len(rows), "emptyRows": empty, "imported": len(out), "withCoords": 0}


def import_fire():
    rows, empty = read_sheet(SRC / "Fire_Station.xlsx")
    out = []
    for i, r in enumerate(rows, 1):
        name = r.get("NAME", "")
        if not name:
            continue
        ward = r.get("WARD_ID", "")
        out.append(
            facility(
                id_=f"fire-{i:03d}",
                name=name,
                address=r.get("LOCATION", ""),
                phone=r.get("FS_CONTACT", ""),
                zone=f"Ward {ward}" if ward else "",
                details=[detail("Ward", "प्रभाग", ward)],
            )
        )
    return out, {"sourceRows": len(rows), "emptyRows": empty, "imported": len(out), "withCoords": 0}


def import_police():
    rows, empty = read_sheet(SRC / "Police_Station.xlsx")
    out = []
    for i, r in enumerate(rows, 1):
        name = r.get("NAME", "")
        if not name:
            continue
        ward = r.get("WARD_ID", "")
        out.append(
            facility(
                id_=f"police-{i:03d}",
                name=name,
                address=r.get("LOCATION", ""),
                phone=r.get("PHONE_NO", ""),
                zone=f"Ward {ward}" if ward else "",
                details=[detail("Ward", "प्रभाग", ward)],
            )
        )
    return out, {"sourceRows": len(rows), "emptyRows": empty, "imported": len(out), "withCoords": 0}


def import_zones():
    rows, empty = read_sheet(SRC / "Zone Office.xlsx")
    out = []
    for i, r in enumerate(rows, 1):
        zone_no = r.get("Zone_No", "")
        if not zone_no:
            continue
        name = f"Zone Office {zone_no}"
        out.append(
            facility(
                id_=f"zone-{i:03d}",
                name=name,
                address=r.get("Address", ""),
                zone=f"Zone {zone_no}",
                details=[detail("Zone", "झोन", zone_no)],
            )
        )
    return out, {"sourceRows": len(rows), "emptyRows": empty, "imported": len(out), "withCoords": 0}


def import_hoardings():
    rows, empty = read_sheet(SRC / "Hoardings.xlsx")
    out = []
    with_coords = 0
    blank_agency = 0
    for i, r in enumerate(rows, 1):
        agency = r.get("Agency_nam", "")
        location = r.get("Location", "")
        licence = r.get("Licence_No", "")
        if not agency:
            blank_agency += 1
        if agency and location:
            name = f"{agency} — {location}"
        elif agency:
            name = agency
        elif location:
            name = f"Hoarding at {location}"
        elif licence:
            name = f"Hoarding licence {licence}"
        else:
            continue
        size = r.get("Size", "")
        size2 = r.get("F5", "")
        size_label = ""
        if size and size2:
            size_label = f"{size} × {size2}"
        elif size:
            size_label = size
        lat, lng = r.get("Lat", ""), r.get("Long", "")
        rec = facility(
            id_=f"hoarding-{i:03d}",
            name=name,
            address=location,
            latitude=lat,
            longitude=lng,
            details=[
                detail("Licence no.", "परवाना क्र.", licence),
                detail("Agency", "एजन्सी", agency),
                detail("Size", "आकार", size_label),
            ],
        )
        if rec["latitude"] and rec["longitude"]:
            with_coords += 1
        out.append(rec)
    return out, {
        "sourceRows": len(rows),
        "emptyRows": empty,
        "imported": len(out),
        "withCoords": with_coords,
        "blankAgencyNames": blank_agency,
        "withoutCoords": len(out) - with_coords,
    }


def import_hospitals():
    rows, empty = read_sheet(SRC / "Hospitals.xlsx")
    out = []
    name_counts = {}
    for r in rows:
        n = r.get("NAME", "")
        name_counts[n] = name_counts.get(n, 0) + 1
    duplicate_names = sum(1 for n, c in name_counts.items() if n and c > 1)

    for i, r in enumerate(rows, 1):
        name = r.get("NAME", "")
        if not name:
            continue
        ward = r.get("WARD_ID", "")
        # EMERGENCY_ in this GIS export stores contact numbers (not a boolean).
        phone = r.get("EMERGENCY_", "")
        # Clean obvious zero placeholders that are not real contacts
        if phone in {"0", "0.0"}:
            phone = ""
        details = [
            detail("Type", "प्रकार", r.get("HOSP_TYPE", "")),
            detail("Category", "वर्ग", r.get("HOSP_CAT", "")),
            detail("Ownership", "मालकी", r.get("OWN_TYPE", "")),
            detail("Ambulance", "रुग्णवाहिका", r.get("AMBU_SERV", "")),
            detail("Beds", "खाटा", r.get("BED_COUNT", "") if r.get("BED_COUNT", "") not in {"0", ""} else ""),
            detail("Blood bank", "रक्तपेढी", r.get("BLOODBANK", "")),
            detail("Mortuary", "मृत्यूगृह", r.get("MORTUARY", "")),
            detail("Recognised by", "मान्यता", r.get("RECOG_BY", "")),
            detail("Police jurisdiction", "पोलीस हद्द", r.get("PS_JURISDI", "")),
            detail("Ward", "प्रभाग", ward),
        ]
        out.append(
            facility(
                id_=f"hospital-{i:03d}",
                name=name,
                address=r.get("LOCATION", ""),
                phone=phone,
                timings=r.get("OPEN_TIME", ""),
                zone=f"Ward {ward}" if ward else "",
                details=details,
            )
        )
    return out, {
        "sourceRows": len(rows),
        "emptyRows": empty,
        "imported": len(out),
        "withCoords": 0,
        "duplicateNameGroups": duplicate_names,
        "note": "HOSP_ID is identical (10) for all rows in source; sequential ids assigned. No lat/lng columns in export.",
    }


def write_json(name: str, records: list):
    path = OUT / name
    path.write_text(json.dumps(records, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return path


def main():
    if not SRC.exists():
        raise SystemExit(f"Missing source folder: {SRC}")

    report = {"categories": {}, "coordinateSummary": {}}

    mapping = [
        ("csmc-schools.json", "schools", import_schools),
        ("fire-stations.json", "fireStations", import_fire),
        ("police-stations.json", "policeStations", import_police),
        ("zone-offices.json", "zoneOffices", import_zones),
        ("hoardings.json", "hoardings", import_hoardings),
        ("csmc-hospitals.json", "hospitals", import_hospitals),
    ]

    total = 0
    with_coords = 0
    for filename, key, fn in mapping:
        records, meta = fn()
        write_json(filename, records)
        coords = sum(1 for r in records if r.get("latitude") and r.get("longitude"))
        meta["withCoords"] = coords
        meta["withoutCoords"] = len(records) - coords
        report["categories"][key] = meta
        total += len(records)
        with_coords += coords
        print(f"wrote {filename}: {len(records)} (coords {coords})")

    # Categories without GIS Excel: empty arrays (no fabricated placeholders)
    for filename in ("phcs.json", "cfcs.json"):
        write_json(filename, [])
        report["categories"][filename.replace(".json", "")] = {
            "imported": 0,
            "note": "No GIS Excel supplied; dummy sample data removed.",
        }
        print(f"wrote {filename}: 0 (no source)")

    # Remove obsolete banner file if present (replaced by hoardings.json)
    old = OUT / "banner-locations.json"
    if old.exists():
        old.unlink()
        report["removed"] = "banner-locations.json"

    report["totals"] = {
        "records": total,
        "withCoordinates": with_coords,
        "withoutCoordinates": total - with_coords,
        "pctWithCoordinates": round(100 * with_coords / total, 1) if total else 0,
    }
    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print("report →", REPORT)
    print("TOTAL", report["totals"])


if __name__ == "__main__":
    main()
